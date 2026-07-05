#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate PPT and syllabus Word for 周敏晖 course application."""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor as PptRGB

BASE_DIR = Path(__file__).parent
PPT_PATH = BASE_DIR / "周敏晖-各论课申报PPT.pptx"
SYLLABUS_PATH = BASE_DIR / "周敏晖-各论课教学大纲.docx"

# Colors
BLUE = PptRGB(0x1A, 0x47, 0x8A)
RED = PptRGB(0xC0, 0x00, 0x00)
GRAY = PptRGB(0x55, 0x55, 0x55)
WHITE = PptRGB(0xFF, 0xFF, 0xFF)


def set_run_font(run, name="微软雅黑", size=18, bold=False, color=None):
    run.font.name = name
    run.font.size = PptPt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # title bar
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9), Inches(1.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    set_run_font(r, size=32, bold=True, color=BLUE)
    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(9), Inches(1.2))
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = subtitle
        set_run_font(r2, size=20, color=GRAY)
    # footer
    ft = slide.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(9), Inches(0.5))
    fp = ft.text_frame.paragraphs[0]
    fp.alignment = PP_ALIGN.CENTER
    fr = fp.add_run()
    fr.text = "课程负责人：周敏晖  |  上海商学院马克思主义学院"
    set_run_font(fr, size=14, color=GRAY)


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(0), Inches(2.5), Inches(10), Inches(1.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(2.7), Inches(9), Inches(1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    set_run_font(r, size=36, bold=True, color=WHITE)


def add_content_slide(prs, title, bullets, sub_bullets=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # header
    hdr = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(9.2), Inches(0.7))
    hp = hdr.text_frame.paragraphs[0]
    hr = hp.add_run()
    hr.text = title
    set_run_font(hr, size=28, bold=True, color=BLUE)
    # line
    line = slide.shapes.add_shape(1, Inches(0.4), Inches(1.0), Inches(9.2), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = RED
    line.line.fill.background()
    # body
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = PptPt(8)
        r = p.add_run()
        r.text = b
        set_run_font(r, size=18, color=GRAY)
    if sub_bullets:
        for sb in sub_bullets:
            p = tf.add_paragraph()
            p.level = 1
            p.space_after = PptPt(4)
            r = p.add_run()
            r.text = sb
            set_run_font(r, size=16, color=GRAY)


def add_table_slide(prs, title, headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    hdr = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(9.2), Inches(0.7))
    hr = hdr.text_frame.paragraphs[0].add_run()
    hr.text = title
    set_run_font(hr, size=26, bold=True, color=BLUE)
    cols = len(headers)
    table_shape = slide.shapes.add_table(len(rows) + 1, cols, Inches(0.3), Inches(1.1), Inches(9.4), Inches(0.4 * (len(rows) + 2)))
    table = table_shape.table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_run_font(r, size=12, bold=True, color=WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=11, color=GRAY)


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "从组织路线到组织再造",
        "习近平基层党建工作思想的继承与发展\n—— 上海商学院各论课申报",
    )

    add_content_slide(
        prs,
        "课程基本信息",
        [
            "课程名称：从组织路线到组织再造：习近平基层党建工作思想的继承与发展",
            "课程性质：思想政治理论课各论（专题选修）",
            "学分/学时：2 学分 / 32 学时",
            "授课对象：全校本科（商科背景为主）",
            "课程负责人：周敏晖（副教授、法学博士）",
            "所在单位：上海商学院马克思主义学院",
        ],
    )

    add_section_slide(prs, "一、课程定位")

    add_content_slide(
        prs,
        "为何开设本课？",
        [
            "回答核心问题：习近平基层党建工作思想如何继承百年组织路线？",
            "区别于习概论：专讲基层党建，以组织学切口深入",
            "区别于一般各论：有田野案例、有原创框架、有实践基地",
            "服务上商特色：商科学生理解组织、治理与服务",
        ],
    )

    add_content_slide(
        prs,
        "课程简介",
        [
            "以负责人在城市社区基层党建与组织再造领域长期研究为基础",
            '提出【谱系—再造—落地】三链分析模型',
            "以上海J街道团队党建为核心样本",
            '结合言子书院开展【行走的思政课】',
            "实现研究特色、理论深度与实践育人统一",
        ],
    )

    add_section_slide(prs, "二、特色研究框架")

    add_content_slide(
        prs,
        '【谱系—再造—落地】三链模型',
        [
            "【谱系链】继承什么 — 百年组织路线演进",
            "  支部建在连上 → 单位制党建 → 社区党建 → 新时代党建引领",
            "【再造链】如何创新 — 以政党为中心的再组织化",
            "  组织弱化 → 组织再造 → 团队党建（趣缘党建）",
            "【落地链】怎样见效 — 组织力转化为治理效能",
            "  组织嵌入 · 理念引领 · 活动聚焦",
            "  主体之维 · 过程之维 · 文化之维",
        ],
    )

    add_table_slide(
        prs,
        "课堂分析工具箱",
        ["工具", "内容", "来源"],
        [
            ["组织力三路径", "组织嵌入、理念引领、活动聚焦", "《理论导刊》2020"],
            ["团队党建三机制", "支部领导团队、党员融入团队、团队凝聚群众", "《当代世界社会主义问题》2019"],
            ["共同体三维度", "主体之维、过程之维、文化之维", "《党政研究》2025"],
        ],
    )

    add_section_slide(prs, "三、负责人研究基础")

    add_table_slide(
        prs,
        "研究成果与课程对照",
        ["成果", "与课程关系"],
        [
            ["专著《城市社区治理的组织再造研究》2022", "课程理论内核"],
            ["《以政党为中心的城市社区再组织化》2020", "谱系链主线"],
            ['《团队党建：城市社区党建工作的新探索》2019', "核心案例"],
            ["《组织力提升路径探析》2020", "落地链工具"],
            ["《构建城市社区共同体：党建何以引领》2025", "三维分析"],
            ['言子书院【行走的思政课】2025', "实践教学"],
        ],
    )

    add_section_slide(prs, "四、教学内容")

    add_table_slide(
        prs,
        "16周教学安排（节选）",
        ["周次", "专题", "三链"],
        [
            ["1-2", "导论、百年组织路线", "谱系链"],
            ["3-6", "组织弱化、组织再造、团队党建", "再造链"],
            ["7-8", "三路径、主体—过程—文化", "落地链"],
            ["9-11", "习论述、自我革命、新兴领域党建", "综合"],
            ["12", "言子书院行走思政课", "实践"],
            ["13-16", "案例汇报、结课答辩", "综合"],
        ],
    )

    add_section_slide(prs, "五、课程创新")

    add_content_slide(
        prs,
        "五大创新点",
        [
            '① 理论框架原创：【谱系—再造—落地】三链模型',
            "② 研究成果转化：专著+CSSCI论文直接进课堂",
            "③ 案例资源独占：上海J街道团队党建长期跟踪",
            "④ 商科院校特色：组织管理视角解读基层党建",
            "⑤ 实践教学基础：言子书院行走思政课已落地",
        ],
    )

    add_content_slide(
        prs,
        "考核方式",
        [
            "平时成绩 20% — 课堂运用三链模型",
            "案例分析作业 30% — 2000字，运用三路径/三维度",
            "行走思政课反思 10% — 1000字实践记录",
            "期末案例研究 40% — 5000字，必用三链模型",
        ],
    )

    add_content_slide(
        prs,
        "预期建设成果",
        [
            "1套完整教案与PPT（含三链模型）",
            "1—2个本土实践教学基地",
            "学生优秀案例报告汇编1册",
            "教改论文或教改项目1项",
            '可复制的【研教一体】各论课范式',
        ],
    )

    add_title_slide(prs, "谢谢！", "敬请各位专家批评指正")

    prs.save(PPT_PATH)
    return PPT_PATH


# ========== Syllabus Word ==========

def set_cell_font(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_h(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_p(doc, text, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    p.paragraph_format.line_spacing = 1.5


def add_tbl(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_font(t.rows[0].cells[i], h, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell_font(t.rows[ri + 1].cells[ci], str(val))
    doc.add_paragraph()


def build_syllabus():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(3.17)
        s.right_margin = Cm(3.17)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("上海商学院课程教学大纲")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("（思想政治理论课各论 · 专题选修）")
    r2.font.size = Pt(14)
    r2.font.name = "宋体"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.add_paragraph()

    add_tbl(doc, ["项目", "内容"], [
        ["课程名称", "从组织路线到组织再造：习近平基层党建工作思想的继承与发展"],
        ["课程代码", "（由教务处填写）"],
        ["课程性质", "思想政治理论课各论 / 专题选修"],
        ["学分", "2"],
        ["总学时", "32"],
        ["理论学时", "24"],
        ["实践学时", "8"],
        ["授课对象", "全校本科学生"],
        ["先修课程", "马克思主义基本原理、中国近现代史纲要（建议）"],
        ["课程负责人", "周敏晖"],
        ["编写日期", "2026年7月"],
    ])

    add_h(doc, "一、课程说明", 1)
    add_p(doc, "本课程是上海商学院马克思主义学院开设的思想政治理论课各论（专题选修）课程，"
          "由周敏晖副教授主讲。课程以负责人在基层党建与组织再造领域的长期研究为基础，"
          "运用【谱系—再造—落地】三链分析模型，系统讲授习近平基层党建工作思想对中国共产党百年组织路线的继承与发展。", indent=True)

    add_h(doc, "二、课程目标", 1)
    for i, g in enumerate([
        "理解中国共产党百年组织路线的演进逻辑；",
        "掌握【谱系—再造—落地】三链分析模型；",
        "运用【组织嵌入、理念引领、活动聚焦】三路径及【主体—过程—文化】三维度分析基层案例；",
        "了解团队党建等组织再造实践及其示范意义；",
        "完成基于上海基层样本的专题调研或案例评析。",
    ], 1):
        add_p(doc, f"（{i}）{g}")

    add_h(doc, "三、教学内容与学时分配", 1)
    weeks = [
        ("1", "2", "导论：为何从【组织】看继承", "总论"),
        ("2", "2", "百年组织路线：从连上到社区", "谱系链"),
        ("3", "2", "单位制解体与【组织弱化】", "谱系链→再造链"),
        ("4", "2", "组织再造：概念、问题与路径", "再造链"),
        ("5", "2", "团队党建：上海J街道案例", "再造链"),
        ("6", "2", "趣缘党建：从地缘到趣缘", "再造链"),
        ("7", "2", "组织力提升：【三路径】", "落地链"),
        ("8", "2", "党建引领：【主体—过程—文化】", "落地链"),
        ("9", "2", "习近平关于基层党建的重要论述", "谱系链"),
        ("10", "2", "从组织路线到自我革命", "谱系链"),
        ("11", "2", "商科视角：楼宇与新兴领域党建", "再造链"),
        ("12", "2", "行走的思政课：言子书院实践", "落地链/实践"),
        ("13", "2", "学员案例汇报（一）", "综合"),
        ("14", "2", "学员案例汇报（二）", "综合"),
        ("15", "2", "三链模型综合与期末指导", "综合"),
        ("16", "2", "结课答辩", "—"),
    ]
    add_tbl(doc, ["序号", "学时", "教学内容", "三链维度"], weeks)

    add_h(doc, "四、教学方法", 1)
    add_p(doc, "本课程采用课堂讲授、案例教学、翻转讨论、行走思政课（现场教学）、"
          "学员案例汇报相结合的方式。强调运用负责人研究成果中的分析工具进行课堂训练。", indent=True)

    add_h(doc, "五、考核方式", 1)
    add_tbl(doc, ["考核项目", "比例", "要求"], [
        ["平时成绩", "20%", "出勤、课堂发言（运用三链模型）"],
        ["案例分析作业", "30%", "2000字，运用三路径或三维度"],
        ["行走思政课反思", "10%", "1000字实践记录"],
        ["期末案例研究", "40%", "5000字，必用三链模型"],
    ])

    add_h(doc, "六、参考教材与文献", 1)
    add_p(doc, "（一）教材与原著")
    refs = [
        "1. 《习近平新时代中国特色社会主义思想学习纲要》，学习出版社、人民出版社。",
        "2. 《习近平著作选读》第一卷、第二卷，人民出版社。",
        "3. 周敏晖：《城市社区治理的组织再造研究》，知识产权出版社，2022。",
    ]
    for ref in refs:
        add_p(doc, ref)
    add_p(doc, "（二）代表性学术论文（课程负责人）")
    papers = [
        "4. 周敏晖：《以政党为中心的城市社区再组织化——以上海市J街道为例》，《长白学刊》2020年第6期。",
        "5. 周敏晖、郝宇青：《团队党建：城市社区党建工作的新探索》，《当代世界社会主义问题》2019年第3期。",
        "6. 周敏晖：《新时代城市社区党组织组织力提升路径探析》，《理论导刊》2020年第6期。",
        "7. 周敏晖、郝宇青：《构建城市社区共同体：党建何以引领》，《党政研究》2025年第2期。",
    ]
    for p in papers:
        add_p(doc, p)

    add_h(doc, "七、课程特色", 1)
    add_p(doc, "本课程以【谱系—再造—落地】三链模型为原创分析框架，以负责人专著及CSSCI系列论文为内容支撑，"
          "以上海J街道团队党建为核心案例，以言子书院行走思政课为实践载体，"
          "体现上海商学院商科背景与【研教一体】建设方向。", indent=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("编写人：周敏晖\n审核人：____________\n日期：2026年7月")
    run.font.size = Pt(12)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.save(SYLLABUS_PATH)
    return SYLLABUS_PATH


if __name__ == "__main__":
    p1 = build_ppt()
    p2 = build_syllabus()
    print(f"PPT: {p1}")
    print(f"Syllabus: {p2}")
