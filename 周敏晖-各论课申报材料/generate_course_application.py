#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Word application document for 周敏晖 course."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUTPUT_PATH = "/workspace/周敏晖-各论课申报书-从组织路线到组织再造.docx"


def set_cell_font(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_title(doc, text, level=0):
    if level == 0:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(22)
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    else:
        doc.add_heading(text, level=level)


def add_body(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(12)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    p.paragraph_format.line_spacing = 1.5
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_font(table.rows[0].cells[i], h, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_font(table.rows[r_idx + 1].cells[c_idx], str(val))
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def build_document():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    add_title(doc, "上海商学院一流本科课程（各论课）申报书")
    add_body(doc, "（思想政治理论课专题选修课程）", bold=False)
    doc.add_paragraph()

    # ========== 一、基本信息 ==========
    add_title(doc, "一、课程基本信息", level=1)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["课程名称", "从组织路线到组织再造：习近平基层党建工作思想的继承与发展"],
            ["课程性质", "思想政治理论课各论（专题选修）/ 校级一流本科课程"],
            ["学分/学时", "2 学分 / 32 学时"],
            ["授课对象", "全校本科学生（以商科类专业为主）"],
            ["课程负责人", "周敏晖"],
            ["所在单位", "上海商学院马克思主义学院"],
            ["负责人职务", "副教授、毛泽东思想和中国特色社会主义理论体系概论教研室主任"],
            ["负责人职称/学历", "副教授 / 法学博士"],
            ["学术兼职", "华东师范大学新兴领域与党建研究中心研究员"],
            ["研究方向", "马克思主义中国化、基层党建、基层治理"],
        ],
        col_widths=[4, 12],
    )

    # ========== 二、课程简介 ==========
    add_title(doc, "二、课程简介", level=1)
    add_body(
        doc,
        "本课程由上海商学院马克思主义学院周敏晖副教授主讲。课程以负责人在城市社区基层党建与组织再造领域长期研究为基础，"
        "提出“谱系—再造—落地”三链分析模型，系统讲授习近平基层党建工作思想对中国共产党百年组织路线的继承与发展。",
        indent=True,
    )
    add_body(
        doc,
        "主要教学内容包括：百年组织路线演进、单位制解体后的组织弱化与组织再造、团队党建等组织创新形态、"
        "组织力提升“三路径”（组织嵌入、理念引领、活动聚焦）、党建引领“主体—过程—文化”三维分析等。"
        "课程以上海基层实践样本为核心案例，结合言子书院等本土资源开展“行走的思政课”实践教学。",
        indent=True,
    )
    add_body(
        doc,
        "本课程面向全校本科生开放，尤其适合商科背景学生从组织与治理视角理解新时代党的建设，"
        "实现研究特色、理论深度与实践育人的有机统一，与现有《毛泽东思想和中国特色社会主义理论体系概论》"
        "《习近平新时代中国特色社会主义思想概论》等必修课形成互补，不重复、能深化。",
        indent=True,
    )

    # ========== 三、课程目标 ==========
    add_title(doc, "三、课程目标", level=1)
    goals = [
        "说清中国共产党百年组织路线的演进逻辑及其与习近平基层党建工作思想的关系；",
        "运用“谱系—再造—落地”三链模型，分析一项基层党建创新案例；",
        "掌握“组织嵌入—理念引领—活动聚焦”及“主体—过程—文化”等分析工具；",
        "理解团队党建等组织再造实践对新时代基层党建的示范意义；",
        "完成一份基于上海基层样本的专题调研报告或案例评析。",
    ]
    for i, g in enumerate(goals, 1):
        add_body(doc, f"{i}. {g}")

    # ========== 四、研究框架 ==========
    add_title(doc, "四、特色研究框架：“谱系—再造—落地”三链模型", level=1)
    add_body(
        doc,
        "本课程区别于一般“习思想概论式”各论课，以组织学视角构建原创分析框架，简称“三链模型”：",
        indent=True,
    )
    add_table(
        doc,
        ["维度", "核心命题", "主要内容"],
        [
            ["谱系链（继承什么）", "组织路线百年演进", "从“支部建在连上”到单位制党建、社区党建，再到新时代党建引领基层治理"],
            ["再造链（如何创新）", "以政党为中心的再组织化", "组织弱化—组织再造；团队党建；从地缘到趣缘的结构创新"],
            ["落地链（怎样见效）", "组织力转化为治理效能", "组织嵌入、理念引领、活动聚焦；主体—过程—文化三维分析"],
        ],
        col_widths=[3.5, 4, 8.5],
    )

    add_body(doc, "课堂分析工具箱（均源于负责人研究成果）：", bold=True)
    add_table(
        doc,
        ["工具名称", "具体内容", "文献来源"],
        [
            ["组织力三路径", "组织嵌入、理念引领、活动聚焦", "周敏晖：《新时代城市社区党组织组织力提升路径探析》，《理论导刊》2020年第6期"],
            ["团队党建三机制", "支部领导团队、党员融入团队、团队凝聚群众", "周敏晖、郝宇青：《“团队党建”：城市社区党建工作的新探索》，《当代世界社会主义问题》2019年第3期"],
            ["共同体三维度", "主体之维、过程之维、文化之维", "周敏晖、郝宇青：《构建城市社区共同体：党建何以引领》，《党政研究》2025年第2期"],
        ],
        col_widths=[3, 5.5, 7.5],
    )

    # ========== 五、教学内容 ==========
    add_title(doc, "五、教学内容与学时安排（16周，32学时）", level=1)
    weeks = [
        ("1", "导论：为何从“组织”看继承", "总论", "专著导论问题意识"),
        ("2", "百年组织路线：从连上到社区", "谱系链", "《以政党为中心的城市社区再组织化》"),
        ("3", "单位制解体与“组织弱化”", "谱系链→再造链", "专著第1—2章逻辑"),
        ("4", "组织再造：概念、问题与路径", "再造链", "专著核心章节"),
        ("5", "团队党建：上海J街道案例", "再造链", "2019年CSSCI论文"),
        ("6", "趣缘党建：从地缘到趣缘", "再造链", "田野案例教学"),
        ("7", "组织力提升：“三路径”", "落地链", "2020年《理论导刊》"),
        ("8", "党建引领：“主体—过程—文化”", "落地链", "2025年《党政研究》"),
        ("9", "习近平关于基层党建的重要论述", "谱系链", "原著导读+研究解读"),
        ("10", "从组织路线到自我革命", "谱系链", "与组织再造衔接"),
        ("11", "商科视角：楼宇、商圈与新兴领域党建", "再造链", "新兴领域与党建研究中心成果"),
        ("12", "行走的思政课：言子书院实践", "落地链", "2025年上商实践教学"),
        ("13", "学员案例汇报（一）", "综合", "学生调研"),
        ("14", "学员案例汇报（二）", "综合", "学生调研"),
        ("15", "三链模型综合与期末指导", "综合", "研究总结"),
        ("16", "结课答辩", "—", "—"),
    ]
    add_table(
        doc,
        ["周次", "专题", "三链维度", "负责人成果融入"],
        weeks,
        col_widths=[1.2, 5.5, 2.5, 7],
    )

    add_body(
        doc,
        "实践教学：第12周开展言子书院“行走的思政课”（负责人已有实践基础）；"
        "可选拓展闵行区江川路街道团队党建参访（负责人长期研究样本点）。实践学时约8学时，占比25%。",
        indent=True,
    )

    # ========== 六、考核方案 ==========
    add_title(doc, "六、考核方式与成绩构成", level=1)
    add_table(
        doc,
        ["考核项目", "比例", "说明"],
        [
            ["平时成绩", "20%", "出勤、课堂发言（运用三链模型）"],
            ["案例分析作业", "30%", "运用“三路径”或“三维度”分析基层案例（2000字）"],
            ["行走思政课反思", "10%", "言子书院或参访记录（1000字）"],
            ["期末案例研究", "40%", "5000字，必须运用三链模型，鼓励引用上海基层样本"],
        ],
        col_widths=[3.5, 2, 10.5],
    )

    # ========== 七、创新点 ==========
    add_title(doc, "七、课程创新点", level=1)
    innovations = [
        ("理论框架原创", "课程负责人基于长期田野研究，提出“谱系—再造—落地”三链模型，以组织学视角阐释习近平基层党建工作思想对百年组织路线的继承与发展，区别于一般概论式各论课。"),
        ("研究成果直接转化", "负责人出版专著《城市社区治理的组织再造研究》，在《社会科学》《当代世界社会主义问题》《理论导刊》《党政研究》等刊发表系列论文，课程内容即研究成果的课堂转化，实现“研教一体”。"),
        ("案例资源不可替代", "以上海闵行区江川路街道“团队党建”为核心样本，形成从地缘到趣缘、从组织弱化到组织再造的完整案例链，具有长期跟踪研究基础。"),
        ("商科院校特色鲜明", "立足上海商学院办学定位，从组织管理、平台协同、服务导向等角度解读基层党建，增强商科学生对“组织路线”“基层治理”的理解力与职业伦理感。"),
        ("实践教学已有基础", "负责人已开展言子书院“行走的思政课”，将“两个结合”与基层传播组织逻辑相结合，具备可复制、可推广的实践教学模式。"),
    ]
    for i, (title, content) in enumerate(innovations, 1):
        add_body(doc, f"创新点{i}：{title}", bold=True)
        add_body(doc, content, indent=True)

    # ========== 八、负责人研究基础 ==========
    add_title(doc, "八、课程负责人研究基础与成果对照", level=1)
    add_body(
        doc,
        "周敏晖，女，1985年9月生，中共党员，法学博士，上海商学院马克思主义学院副教授，"
        "毛泽东思想和中国特色社会主义理论体系概论教研室主任，华东师范大学新兴领域与党建研究中心研究员。"
        "主持和参与省部级以上课题十余项，出版专著一部，在CSSCI期刊发表论文十余篇，在《解放日报》《上海宣传通讯》等主流媒体发表论文多篇。",
        indent=True,
    )
    add_table(
        doc,
        ["成果类型", "成果名称/出处", "与课程关系"],
        [
            ["专著", "《城市社区治理的组织再造研究》，知识产权出版社，2022", "课程理论内核，第3—4讲"],
            ["CSSCI论文", "《以政党为中心的城市社区再组织化——以上海市J街道为例》，《长白学刊》2020年第6期", "谱系链主线，第2讲"],
            ["CSSCI论文", "《“团队党建”：城市社区党建工作的新探索》，《当代世界社会主义问题》2019年第3期", "核心案例，第5—6讲"],
            ["CSSCI论文", "《新时代城市社区党组织组织力提升路径探析》，《理论导刊》2020年第6期", "落地链工具，第7讲"],
            ["CSSCI论文", "《构建城市社区共同体：党建何以引领》，《党政研究》2025年第2期", "三维分析，第8讲"],
            ["实践教学", "言子书院“行走的思政课”（2025，上海商学院官网报道）", "实践教学，第12讲"],
        ],
        col_widths=[2.5, 7.5, 6],
    )

    # ========== 九、与现有课程关系 ==========
    add_title(doc, "九、与学院现有思政课程的关系", level=1)
    add_table(
        doc,
        ["现有课程", "本课程与之关系"],
        [
            ["《毛泽东思想和中国特色社会主义理论体系概论》", "本课从组织路线角度深化毛概中关于党的建设、基层组织的相关内容，形成专题延伸"],
            ["《习近平新时代中国特色社会主义思想概论》", "本课专讲习思想中基层党建板块，以组织再造为切口，不重复概论体系"],
            ["《思想道德与法治》", "本课侧重组织伦理与基层治理，可与德法课形成“价值—组织—实践”互补"],
        ],
        col_widths=[6, 10],
    )

    # ========== 十、预期成果 ==========
    add_title(doc, "十、预期建设成果", level=1)
    expected = [
        "形成1套基于“三链模型”的专题课程完整教案与PPT；",
        "建设1—2个上海本土实践教学基地（言子书院、基层社区样本点）；",
        "产出学生优秀案例研究报告汇编1册；",
        "发表教改论文1篇，或申报校级/市级教改项目1项；",
        "为马克思主义学院各论课建设提供可复制的“研教一体”范式。",
    ]
    for i, e in enumerate(expected, 1):
        add_body(doc, f"{i}. {e}")

    # ========== 十一、承诺 ==========
    add_title(doc, "十一、负责人承诺", level=1)
    add_body(
        doc,
        "本人承诺：本申报书所填内容真实准确；如获立项，将严格按照学校一流本科课程建设要求完成课程建设任务，"
        "保证课程质量与教学投入，接受学校与学院的教学检查与评估。",
        indent=True,
    )
    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("课程负责人（签字）：______________")
    run.font.size = Pt(12)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run2 = p2.add_run("申报日期：2026年____月____日")
    run2.font.size = Pt(12)
    run2.font.name = "宋体"
    run2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run3 = p3.add_run("上海商学院马克思主义学院（盖章）")
    run3.font.size = Pt(12)
    run3.font.name = "宋体"
    run3._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(f"Generated: {path}")
